"""Parse and normalize DipTrace bill-of-material files."""

from __future__ import annotations

import csv
import io
import re
from dataclasses import asdict, dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import BinaryIO, Iterable


class BomParseError(ValueError):
    """Raised when an uploaded BOM cannot be parsed safely."""


@dataclass(frozen=True)
class BomRow:
    """A normalized DipTrace BOM line."""

    row: int
    designators: str
    footprint: str
    comment: str
    quantity: str
    jlcpcb_part: str

    def as_dict(self) -> dict:
        return asdict(self)


HEADER_ALIASES = {
    "designators": {"designator", "designators", "reference", "references", "refdes"},
    "footprint": {"footprint", "package", "pattern"},
    "comment": {"comment", "value", "description", "part", "component"},
    "quantity": {"quantity", "qty", "count"},
    "jlcpcb_part": {
        "jlcpcb part #",
        "jlcpcb part#",
        "jlcpcb part",
        "jlc part #",
        "jlc part",
        "lcsc part #",
        "lcsc part#",
        "lcsc part",
        "supplier part",
        "supplier sku",
    },
}


def parse_bom(upload: BinaryIO, filename: str) -> list[BomRow]:
    """Parse CSV or XLSX content and return consolidated normalized rows."""
    suffix = Path(filename or "").suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        records = _xlsx_records(upload)
    elif suffix in {".csv", ".txt"}:
        records = _csv_records(upload)
    else:
        raise BomParseError("Upload a DipTrace CSV or XLSX file")

    return normalize_records(records)


def normalize_records(records: Iterable[dict]) -> list[BomRow]:
    """Normalize records and consolidate identical grouped component rows."""
    records = list(records)
    if not records:
        raise BomParseError("The uploaded file is empty")

    header_map = _header_map(records[0].keys())
    missing = [name for name in ("designators", "quantity") if name not in header_map]
    if missing:
        raise BomParseError(f"Missing required column(s): {', '.join(missing)}")

    grouped: dict[tuple[str, str, str], dict] = {}
    order: list[tuple[str, str, str]] = []

    for input_row, record in enumerate(records, start=2):
        designators = _clean(record.get(header_map["designators"], ""))
        footprint = _clean(record.get(header_map.get("footprint", ""), ""))
        comment = _clean(record.get(header_map.get("comment", ""), ""))
        supplier_code = _normalize_jlc_code(
            record.get(header_map.get("jlcpcb_part", ""), "")
        )
        quantity_value = _clean(record.get(header_map["quantity"], ""))

        # DipTrace commonly emits a blank final row containing only the total quantity.
        if not any((designators, footprint, comment, supplier_code)):
            continue
        if not designators:
            raise BomParseError(f"Row {input_row}: designator is required")

        try:
            quantity = Decimal(quantity_value)
        except (InvalidOperation, TypeError):
            raise BomParseError(f"Row {input_row}: invalid quantity {quantity_value!r}") from None

        if quantity <= 0:
            raise BomParseError(f"Row {input_row}: quantity must be greater than zero")

        key = (supplier_code.casefold(), footprint.casefold(), comment.casefold())
        if key not in grouped:
            grouped[key] = {
                "row": input_row,
                "designators": [],
                "footprint": footprint,
                "comment": comment,
                "quantity": Decimal("0"),
                "jlcpcb_part": supplier_code,
            }
            order.append(key)

        grouped[key]["quantity"] += quantity
        grouped[key]["designators"].extend(_split_designators(designators))

    if not order:
        raise BomParseError("No component rows were found in the uploaded file")

    rows = []
    for output_row, key in enumerate(order, start=1):
        item = grouped[key]
        rows.append(
            BomRow(
                row=output_row,
                designators=", ".join(_deduplicate(item["designators"])),
                footprint=item["footprint"],
                comment=item["comment"],
                quantity=_decimal_string(item["quantity"]),
                jlcpcb_part=item["jlcpcb_part"],
            )
        )

    return rows


def _csv_records(upload: BinaryIO) -> list[dict]:
    raw = upload.read()
    if isinstance(raw, str):
        text = raw
    else:
        for encoding in ("utf-8-sig", "utf-16", "cp1252"):
            try:
                text = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise BomParseError("Could not decode the CSV file")

    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel

    return list(csv.DictReader(io.StringIO(text), dialect=dialect))


def _xlsx_records(upload: BinaryIO) -> list[dict]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - declared runtime dependency
        raise BomParseError("XLSX support requires openpyxl") from exc

    try:
        workbook = load_workbook(upload, read_only=True, data_only=True)
        sheet = workbook.active
        values = sheet.iter_rows(values_only=True)
        headers = next(values)
    except Exception as exc:
        raise BomParseError(f"Could not read the Excel workbook: {exc}") from exc

    names = [_clean(value) for value in headers]
    return [dict(zip(names, row, strict=False)) for row in values]


def _header_map(headers: Iterable[str]) -> dict[str, str]:
    normalized = {_clean(header).casefold(): header for header in headers if header is not None}
    result = {}
    for canonical, aliases in HEADER_ALIASES.items():
        for alias in aliases:
            if alias in normalized:
                result[canonical] = normalized[alias]
                break
    return result


def _normalize_jlc_code(value) -> str:
    value = _clean(value).upper().replace(" ", "")
    if not value:
        return ""
    match = re.fullmatch(r"C(\d+)", value)
    return f"C{match.group(1)}" if match else value


def _split_designators(value: str) -> list[str]:
    return [part.strip() for part in re.split(r"[,;\s]+", value) if part.strip()]


def _deduplicate(values: Iterable[str]) -> list[str]:
    result = []
    seen = set()
    for value in values:
        key = value.casefold()
        if key not in seen:
            seen.add(key)
            result.append(value)
    return result


def _decimal_string(value: Decimal) -> str:
    return format(value.normalize(), "f")


def _clean(value) -> str:
    if value is None:
        return ""
    return " ".join(str(value).replace("\ufeff", "").strip().split())

